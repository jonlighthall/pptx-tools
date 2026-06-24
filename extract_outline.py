#!/usr/bin/env python
"""Extract a simple markdown outline from a PowerPoint .pptx file.

Usage:
    python extract_outline.py input.pptx [output.md]

- One top-level heading per slide: "# Slide N: <title>"
- Bullet lines for other text boxes on the slide.
- If no output path is given, writes "<stem>.export.md" (e.g. talk_r3.pptx ->
  talk_r3.export.md). The ".export.md" suffix marks the file as a machine-
  generated export tied to its source .pptx, so it never collides with hand-
  written working drafts that share the same stem.

Requires: python-pptx
"""

import sys
from pathlib import Path

from pptx import Presentation


def iter_slide_text(slide):
    """Yield (is_title, text) pairs for the given slide.

    - The slide title (if present) is returned first with is_title=True.
    - Other text-containing shapes are returned with is_title=False.
    """
    title_shape = slide.shapes.title
    title_text = None
    if title_shape is not None and getattr(title_shape, "has_text_frame", False):
        title_text = title_shape.text.strip()
        if title_text:
            yield True, title_text

    # Collect other text shapes, skipping the title shape
    for shape in slide.shapes:
        if shape is title_shape:
            continue
        if not getattr(shape, "has_text_frame", False):
            continue
        text = shape.text.strip()
        if not text:
            continue
        # Avoid duplicating the title text
        if title_text and text == title_text:
            continue
        # Split into lines to avoid huge blocks
        for line in text.splitlines():
            line = line.strip()
            if line:
                yield False, line


def main(argv: list[str]) -> int:
    if len(argv) not in (2, 3):
        print("Usage: python extract_outline.py input.pptx [output.md]", file=sys.stderr)
        return 1

    in_path = Path(argv[1])
    # Default output: "<stem>.export.md" (NOT "<stem>.md").
    # The ".export.md" suffix marks the file as a machine-generated export tied to
    # its source .pptx, keeping it distinct from hand-written working drafts and
    # preventing accidental clobbering of drafts that share the stem.
    out_path = Path(argv[2]) if len(argv) == 3 else in_path.with_suffix(".export.md")

    if not in_path.is_file():
        print(f"Input file not found: {in_path}", file=sys.stderr)
        return 1

    prs = Presentation(in_path)

    lines: list[str] = []

    for idx, slide in enumerate(prs.slides, start=1):
        title = None
        body_lines: list[str] = []

        for is_title, text in iter_slide_text(slide):
            if is_title:
                # First title wins
                if title is None:
                    title = text
            else:
                body_lines.append(text)

        # Fallback title if none was found
        if title:
            heading = f"# Slide {idx}: {title}"
        else:
            heading = f"# Slide {idx}"

        lines.append(heading)

        for bl in body_lines:
            lines.append(f"- {bl}")

        # Blank line between slides
        lines.append("")

    out_path.write_text("\n".join(lines), encoding="utf-8")
    print(f"Wrote outline to {out_path}")
    return 0


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit(main(sys.argv))
